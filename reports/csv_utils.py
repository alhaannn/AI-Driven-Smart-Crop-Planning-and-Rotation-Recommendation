"""
CSV and Excel export utilities
"""
import csv
from io import BytesIO
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime


def export_fields_to_csv():
    """Export all fields to CSV"""
    from fields.models import Field
    
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="fields_{datetime.now().strftime("%Y%m%d")}.csv"'
    
    writer = csv.writer(response)
    writer.writerow(['Field Name', 'Location', 'Size (acres)', 'Soil Type', 'Irrigation Type', 'Created Date'])
    
    fields = Field.objects.all()
    for field in fields:
        writer.writerow([
            field.name,
            field.location,
            field.size_acres,
            field.soil_type,
            field.irrigation_type,
            field.created_at.strftime('%Y-%m-%d')
        ])
    
    return response


def export_crop_history_to_csv(field=None):
    """Export crop history to CSV"""
    from fields.models import CropHistory
    
    response = HttpResponse(content_type='text/csv')
    filename = f'crop_history_{field.name if field else "all"}_{datetime.now().strftime("%Y%m%d")}.csv'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    writer = csv.writer(response)
    writer.writerow(['Field', 'Crop Name', 'Variety', 'Planting Date', 'Harvest Date', 'Yield (kg/acre)', 'Notes'])
    
    crops = CropHistory.objects.all() if not field else field.crop_histories.all()
    for crop in crops:
        writer.writerow([
            crop.field.name,
            crop.crop_name,
            crop.variety or 'N/A',
            crop.planting_date.strftime('%Y-%m-%d'),
            crop.harvest_date.strftime('%Y-%m-%d'),
            crop.yield_amount,
            crop.notes or ''
        ])
    
    return response


def export_soil_analysis_to_csv(field=None):
    """Export soil analysis to CSV"""
    from soil.models import SoilAnalysis
    
    response = HttpResponse(content_type='text/csv')
    filename = f'soil_analysis_{field.name if field else "all"}_{datetime.now().strftime("%Y%m%d")}.csv'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    writer = csv.writer(response)
    writer.writerow(['Field', 'Analysis Date', 'pH Level', 'Nitrogen (ppm)', 'Phosphorus (ppm)', 
                     'Potassium (ppm)', 'Organic Matter %', 'Notes'])
    
    analyses = SoilAnalysis.objects.all() if not field else field.soil_analyses.all()
    for analysis in analyses:
        writer.writerow([
            analysis.field.name,
            analysis.analysis_date.strftime('%Y-%m-%d'),
            analysis.ph_level,
            analysis.nitrogen_ppm,
            analysis.phosphorus_ppm,
            analysis.potassium_ppm,
            analysis.organic_matter,
            analysis.notes or ''
        ])
    
    return response


def export_financial_records_to_csv(field=None, start_date=None, end_date=None):
    """Export financial records to CSV"""
    from analytics.models import FinancialRecord
    
    response = HttpResponse(content_type='text/csv')
    filename = f'financial_{field.name if field else "all"}_{datetime.now().strftime("%Y%m%d")}.csv'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    writer = csv.writer(response)
    writer.writerow(['Field', 'Date', 'Type', 'Category', 'Amount (₹)', 'Description', 'Payment Method'])
    
    records = FinancialRecord.objects.all() if not field else FinancialRecord.objects.filter(field=field)
    
    if start_date:
        records = records.filter(date__gte=start_date)
    if end_date:
        records = records.filter(date__lte=end_date)
    
    for record in records:
        writer.writerow([
            record.field.name if record.field else 'N/A',
            record.date.strftime('%Y-%m-%d'),
            record.get_transaction_type_display(),
            record.get_category_display(),
            record.amount,
            record.description or '',
            record.get_payment_method_display() if record.payment_method else 'N/A'
        ])
    
    return response


def export_fields_to_excel():
    """Export all fields to Excel with professional formatting"""
    from fields.models import Field
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Fields"
    
    # Define styles
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Headers
    headers = ['Field Name', 'Location', 'Size (acres)', 'Soil Type', 'Irrigation Type', 'Created Date']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # Data
    fields = Field.objects.all()
    for row_num, field in enumerate(fields, 2):
        ws.cell(row=row_num, column=1, value=field.name).border = border
        ws.cell(row=row_num, column=2, value=field.location).border = border
        ws.cell(row=row_num, column=3, value=float(field.size_acres)).border = border
        ws.cell(row=row_num, column=4, value=field.soil_type).border = border
        ws.cell(row=row_num, column=5, value=field.irrigation_type).border = border
        ws.cell(row=row_num, column=6, value=field.created_at.strftime('%Y-%m-%d')).border = border
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 15
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Create response
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="fields_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    
    return response


def export_financial_to_excel(field=None, start_date=None, end_date=None):
    """Export financial records to Excel with summary"""
    from analytics.models import FinancialRecord
    
   # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Financial Records"
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Title
    ws['A1'] = 'Financial Report'
    ws['A1'].font = Font(bold=True, size=16)
    ws.merge_cells('A1:G1')
    ws['A1'].alignment = Alignment(horizontal='center')
    
    # Headers
    headers = ['Field', 'Date', 'Type', 'Category', 'Amount (₹)', 'Description', 'Payment Method']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # Data
    records = FinancialRecord.objects.all() if not field else FinancialRecord.objects.filter(field=field)
    
    if start_date:
        records = records.filter(date__gte=start_date)
    if end_date:
        records = records.filter(date__lte=end_date)
    
    row_num = 4
    total_income = 0
    total_expense = 0
    
    for record in records:
        ws.cell(row=row_num, column=1, value=record.field.name if record.field else 'N/A').border = border
        ws.cell(row=row_num, column=2, value=record.date.strftime('%Y-%m-%d')).border = border
        ws.cell(row=row_num, column=3, value=record.get_transaction_type_display()).border = border
        ws.cell(row=row_num, column=4, value=record.get_category_display()).border = border
        ws.cell(row=row_num, column=5, value=float(record.amount)).border = border
        ws.cell(row=row_num, column=6, value=record.description or '').border = border
        ws.cell(row=row_num, column=7, value=record.get_payment_method_display() if record.payment_method else 'N/A').border = border
        
        if record.transaction_type == 'income':
            total_income += float(record.amount)
        else:
            total_expense += float(record.amount)
        
        row_num += 1
    
    # Summary
    row_num += 2
    ws.cell(row=row_num, column=4, value='Total Income:').font = Font(bold=True)
    ws.cell(row=row_num, column=5, value=total_income).font = Font(bold=True)
    
    row_num += 1
    ws.cell(row=row_num, column=4, value='Total Expenses:').font = Font(bold=True)
    ws.cell(row=row_num, column=5, value=total_expense).font = Font(bold=True)
    
    row_num += 1
    ws.cell(row=row_num, column=4, value='Net Profit/Loss:').font = Font(bold=True, color="FF0000" if (total_income - total_expense) < 0 else "008000")
    ws.cell(row=row_num, column=5, value=total_income - total_expense).font = Font(bold=True, color="FF0000" if (total_income - total_expense) < 0 else "008000")
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 35
    ws.column_dimensions['G'].width = 18
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Create response
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'financial_{field.name if field else "all"}_{datetime.now().strftime("%Y%m%d")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    return response
